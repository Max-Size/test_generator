"""Извлечение содержимого из загруженных файлов в Markdown: pdf / docx / txt / md.

PDF обрабатывается послойно:
1. `pymupdf4llm` вытаскивает текст, заголовки и таблицы в Markdown — быстро и без LLM;
2. «тяжёлые» страницы (картинки, графики, формулы, почти нет текстового слоя)
   дорисовываются vision-моделью: страница рендерится в PNG и распознаётся в Markdown
   с формулами в LaTeX и текстовым описанием иллюстраций. Распознавание идёт параллельно
   (по одной странице на запрос, не более `vision_concurrency` одновременно).

Результат извлечения кэшируется на диске по SHA-256 файла — повторная обработка того же
документа не тратит ни время, ни токены.
"""
import asyncio
import hashlib
import io
from pathlib import Path

import pymupdf
import pymupdf4llm
from docx import Document
from fastapi import HTTPException

from .config import settings
from .llm import PageTranscription, transcribe_page
from .obs import log
from .schemas import TruncatedFile
from .text_clean import clean_markdown, page_has_symbol_font, page_text_from_dict


def _digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _cache_file(data: bytes) -> Path:
    return Path(settings.ingest_cache_dir) / f"{_digest(data)}.md"


def _diagrams_dir(data: bytes) -> Path:
    return Path(settings.ingest_cache_dir) / f"{_digest(data)}.diagrams"


def _crop_diagram(pix: "pymupdf.Pixmap", bbox: list[float]) -> bytes | None:
    """Вырезать bbox (доли 0..1) из пиксмапа; None, если площадь после клампа нулевая."""
    if len(bbox) != 4:
        return None
    x0, y0, x1, y1 = (min(max(v, 0.0), 1.0) for v in bbox)
    left, right = round(min(x0, x1) * pix.width), round(max(x0, x1) * pix.width)
    top, bottom = round(min(y0, y1) * pix.height), round(max(y0, y1) * pix.height)
    if right <= left or bottom <= top:
        return None
    crop = pymupdf.Pixmap(pix.colorspace, pymupdf.IRect(left, top, right, bottom), pix.alpha)
    crop.copy(pix, pymupdf.IRect(left, top, right, bottom))
    return crop.tobytes("png")


def _page_needs_vision(page: "pymupdf.Page", text: str) -> bool:
    if not settings.vision_enabled:
        return False
    return (
        len(text.strip()) < settings.vision_min_text_chars
        or bool(page.get_images())
        or page_has_symbol_font(page)
        or bool(page.get_drawings())
    )


async def _from_pdf(data: bytes) -> tuple[str, dict[str, bytes]]:
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        chunks = pymupdf4llm.to_markdown(doc, page_chunks=True, show_progress=False)
        log.info("pdf: %d страниц", len(chunks))

        # Текстовый слой: страницы с математикой (шрифт Symbol) собираем сами —
        # с ремапом Symbol-глифов и восстановлением индексов; остальные берём из
        # pymupdf4llm (его таблицы/заголовки лучше). Всё чистим от разметочного шума.
        pages_md: list[str] = []
        symbol_pages = 0
        for i, chunk in enumerate(chunks):
            if page_has_symbol_font(doc[i]):
                symbol_pages += 1
                pages_md.append(clean_markdown(page_text_from_dict(doc[i])))
            else:
                pages_md.append(clean_markdown(chunk.get("text") or ""))
        if symbol_pages:
            log.info("pdf: %d страниц с математикой нормализованы без LLM", symbol_pages)

        # какие страницы отдаём в vision-модель (с учётом бюджета на файл)
        heavy: list[int] = []
        for i, md_text in enumerate(pages_md):
            if len(heavy) >= settings.vision_max_pages:
                break
            if _page_needs_vision(doc[i], md_text):
                heavy.append(i)

        diagrams: dict[str, bytes] = {}
        prefix = _digest(data)[:8]
        if heavy:
            sem = asyncio.Semaphore(settings.vision_concurrency)

            async def _recognize(idx: int) -> tuple[int, PageTranscription, "pymupdf.Pixmap"]:
                pix = doc[idx].get_pixmap(dpi=settings.vision_dpi)
                async with sem:
                    return idx, await transcribe_page(pix.tobytes("png"), idx + 1), pix

            log.info("pdf: %d страниц в vision (параллельно по %d)",
                     len(heavy), settings.vision_concurrency)
            for idx, tr, pix in await asyncio.gather(*(_recognize(i) for i in heavy)):
                md = tr.markdown.strip()
                if not md:
                    continue
                log.info("pdf стр. %d: распознана vision-моделью (%d симв.)", idx + 1, len(md))
                for d in tr.diagrams[: settings.vision_max_diagrams_per_page]:
                    png = _crop_diagram(pix, d.bbox)
                    if png is None:
                        log.warning("pdf стр. %d: рисунок %s с вырожденным bbox пропущен",
                                    idx + 1, d.id)
                        continue
                    # id уникален в рамках файла и страницы, а не только страницы
                    new_id = f"{prefix}_p{idx + 1}_{d.id}"
                    md = md.replace(f"diagram_id={d.id}]", f"diagram_id={new_id}]")
                    diagrams[new_id] = png
                pages_md[idx] = md

        return "\n\n".join(p for p in pages_md if p), diagrams
    finally:
        doc.close()


def _from_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    log.info("docx: %d абзацев", len(doc.paragraphs))
    return "\n".join(p.text for p in doc.paragraphs)


def _from_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


async def extract_text(filename: str, data: bytes) -> tuple[str, dict[str, bytes]]:
    name = filename.lower()

    cache = _cache_file(data)
    if cache.exists():
        log.info("извлечение %s взято из кэша (%s)", filename, cache.name)
        ddir = _diagrams_dir(data)
        diagrams = {f.stem: f.read_bytes() for f in ddir.glob("*.png")} if ddir.is_dir() else {}
        return cache.read_text(encoding="utf-8"), diagrams

    diagrams: dict[str, bytes] = {}
    try:
        if name.endswith(".pdf"):
            text, diagrams = await _from_pdf(data)
        elif name.endswith(".docx"):
            text = _from_docx(data)
        elif name.endswith((".txt", ".md")):
            text = _from_text(data)
        else:
            raise HTTPException(415, f"неподдерживаемый формат: {filename} (нужен pdf/docx/txt/md)")
    except HTTPException:
        raise
    except Exception as e:  # noqa: BLE001
        raise HTTPException(422, f"не удалось прочитать {filename}: {e}")

    cache.parent.mkdir(parents=True, exist_ok=True)
    if diagrams:
        ddir = _diagrams_dir(data)
        ddir.mkdir(parents=True, exist_ok=True)
        for did, png in diagrams.items():
            (ddir / f"{did}.png").write_bytes(png)
    cache.write_text(text, encoding="utf-8")  # .md последним: он маркер целостного кэша
    return text, diagrams


def combine(parts: list[tuple[str, str]], limit: int) -> tuple[str, list[TruncatedFile]]:
    """Уложить файлы в промпт целиком, по порядку загрузки, пока укладываются в limit
    символов. Файл, на котором лимит заканчивается, обрезается по последней границе
    абзаца (``\\n\\n``); если в помещающейся части абзаца не нашлось — файл выпадает
    целиком, а не режется посреди слова. Файлы после точки обрезки не попадают в промпт.
    Возвращает (текст, список обрезанных/выпавших файлов).
    """
    chunks: list[str] = []
    truncated: list[TruncatedFile] = []
    used = 0
    cut = False

    for name, text in parts:
        text = text.strip()
        if not text:
            continue
        chunk = f"### {name}\n{text}"

        if cut:
            truncated.append(TruncatedFile(filename=name, reason="dropped"))
            continue

        sep = 2 if chunks else 0  # длина "\n\n" перед этим чанком при финальной склейке
        remaining = limit - used - sep
        if len(chunk) <= remaining:
            chunks.append(chunk)
            used += len(chunk) + sep
            continue

        if remaining <= 0:
            truncated.append(TruncatedFile(filename=name, reason="dropped"))
            cut = True
            continue

        piece = chunk[:remaining]
        boundary = piece.rfind("\n\n")
        if boundary <= 0:
            truncated.append(TruncatedFile(filename=name, reason="dropped"))
        else:
            chunks.append(piece[:boundary])
            truncated.append(TruncatedFile(filename=name, reason="partial"))
        cut = True

    return "\n\n".join(chunks), truncated
